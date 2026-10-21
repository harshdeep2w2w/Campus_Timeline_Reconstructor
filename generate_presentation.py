"""
Campus Network Attack - Timeline Reconstructor
Simple, Easy-to-Understand Academic Presentation Generator
Generates a 17-slide (16:9) presentation in plain, clear English.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_simple_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors (Clean, Modern Dark SOC Theme)
    BG_COLOR = RGBColor(11, 15, 25)          # Deep Navy #0B0F19
    CARD_BG = RGBColor(24, 32, 47)          # Soft Dark Slate #18202F
    CARD_BORDER = RGBColor(45, 60, 90)      # Subtle Border #2D3C5A
    CYAN = RGBColor(56, 189, 248)           # Electric Cyan #38BDF8
    WHITE = RGBColor(248, 250, 252)         # Crisp White #F8FAFC
    LIGHT_GRAY = RGBColor(203, 213, 225)    # Readable Text #CBD5E1
    MUTED = RGBColor(148, 163, 184)         # Subtitles #94A3B8
    RED = RGBColor(239, 68, 68)             # Danger/Alert #EF4444
    AMBER = RGBColor(245, 158, 11)          # Warning/Notice #F59E0B
    GREEN = RGBColor(16, 185, 129)          # Good/Success #10B981
    BLUE_PILL = RGBColor(30, 58, 138)       # Tag #1E3A8A

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title, category="COLLEGE PROJECT • CYBER SECURITY"):
        # Top Small Tag
        tbox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.3))
        tf = tbox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = category.upper()
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = CYAN

        # Main Slide Title
        tbox2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.6))
        tf2 = tbox2.text_frame
        tf2.word_wrap = True
        tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
        p2 = tf2.paragraphs[0]
        p2.text = title
        p2.font.name = "Arial"
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = WHITE

        # Divider Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.fill.background()

    def add_card(slide, left, top, width, height, title="", title_color=CYAN, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)

        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.18), width - Inches(0.4), Inches(0.35))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = title
            p.font.name = "Arial"
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = title_color

        return card

    def add_bullet_points(slide, left, top, width, height, points):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        for idx, pt in enumerate(points):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.space_after = Pt(10)
            p.font.name = "Arial"
            p.font.size = Pt(12)

            # Bullet
            run_b = p.add_run()
            run_b.text = "• "
            run_b.font.bold = True
            run_b.font.color.rgb = CYAN

            if isinstance(pt, tuple):
                bold_part, regular_part = pt
                rb = p.add_run()
                rb.text = bold_part + " "
                rb.font.bold = True
                rb.font.color.rgb = WHITE

                rr = p.add_run()
                rr.text = regular_part
                rr.font.bold = False
                rr.font.color.rgb = LIGHT_GRAY
            else:
                rr = p.add_run()
                rr.text = pt
                rr.font.bold = False
                rr.font.color.rgb = LIGHT_GRAY

    def set_note(slide, text):
        tf = slide.notes_slide.notes_text_frame
        tf.text = text

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    c1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.0), Inches(10.933), Inches(5.5))
    c1.fill.solid()
    c1.fill.fore_color.rgb = RGBColor(16, 23, 40)
    c1.line.color.rgb = RGBColor(37, 99, 235)
    c1.line.width = Pt(1.5)

    # Pill
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.6), Inches(1.5), Inches(3.6), Inches(0.42))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BLUE_PILL
    badge.line.color.rgb = CYAN
    btf = badge.text_frame
    btf.margin_top = Inches(0.06)
    bp = btf.paragraphs[0]
    bp.alignment = PP_ALIGN.CENTER
    bp.text = "B.TECH CSE (CYBER SECURITY)"
    bp.font.name = "Arial"
    bp.font.size = Pt(10)
    bp.font.bold = True
    bp.font.color.rgb = WHITE

    # Title
    tbox = s1.shapes.add_textbox(Inches(1.6), Inches(2.2), Inches(10.1), Inches(1.6))
    tf = tbox.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "Campus Network Attack"
    p1.font.name = "Arial"
    p1.font.size = Pt(38)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    p2 = tf.add_paragraph()
    p2.text = "Timeline Reconstructor"
    p2.font.name = "Arial"
    p2.font.size = Pt(38)
    p2.font.bold = True
    p2.font.color.rgb = CYAN

    # Subtitle in plain english
    subbox = s1.shapes.add_textbox(Inches(1.6), Inches(3.9), Inches(10.1), Inches(0.7))
    stf = subbox.text_frame
    stf.word_wrap = True
    sp = stf.paragraphs[0]
    sp.text = "A Smart Security Tool that Rebuilds Cyber Attack Timelines from Campus Network Logs"
    sp.font.name = "Arial"
    sp.font.size = Pt(15)
    sp.font.color.rgb = LIGHT_GRAY

    # 3 Summary Mini-cards
    items = [
        ("WHAT IT DOES", "Finds hackers & rebuilds the attack story", CYAN),
        ("EASY TO USE", "Visual dashboard with interactive network map", GREEN),
        ("VERIFIED & READY", "Passed all 68 automated software tests", AMBER)
    ]
    for i, (m_title, m_val, m_col) in enumerate(items):
        mx = Inches(1.6 + i * 3.4)
        my = Inches(4.8)
        mc = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, mx, my, Inches(3.2), Inches(1.2))
        mc.fill.solid()
        mc.fill.fore_color.rgb = RGBColor(22, 30, 49)
        mc.line.color.rgb = CARD_BORDER
        mctf = mc.text_frame
        mctf.margin_left = Inches(0.15)
        mctf.margin_top = Inches(0.15)
        mcp1 = mctf.paragraphs[0]
        mcp1.text = m_title
        mcp1.font.name = "Arial"
        mcp1.font.size = Pt(9)
        mcp1.font.bold = True
        mcp1.font.color.rgb = m_col
        mcp2 = mctf.add_paragraph()
        mcp2.text = m_val
        mcp2.font.name = "Arial"
        mcp2.font.size = Pt(11)
        mcp2.font.bold = True
        mcp2.font.color.rgb = WHITE

    set_note(s1,
        "Good morning teacher. Today I am presenting my project called 'Campus Network Attack - Timeline Reconstructor'. "
        "In simple words, this project acts like a digital detective for college campus networks. "
        "When an attack happens, it reads network logs, finds what the hacker did step-by-step, shows the story on an easy timeline, "
        "and draws a visual map of all connected computers."
    )

    # =========================================================================
    # SLIDE 2: WHAT IS THIS PROJECT? (THE BIG PICTURE)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "What is This Project? (The Big Picture)")

    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3), "The Simple Idea Behind the Project", CYAN)
    add_bullet_points(s2, Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.5), [
        ("Like a CCTV for Networks:", "Just like security cameras record people entering buildings, network routers and firewalls record every website visited and computer connection."),
        ("The Problem with Logs:", "These logs are stored in thousands of plain text lines. Reading them by hand is slow and very confusing."),
        ("What Our Software Does:", "It collects all these messy logs, cleans them up, connects the dots, and reconstructs the exact story of how a hacker broke in."),
        ("Who Benefits:", "College IT teams, security teachers, students, and network administrators who need fast answers when something goes wrong.")
    ])

    add_card(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3), "What Can the User Do With It?", GREEN)
    add_bullet_points(s2, Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.5), [
        ("Upload Any Log File:", "Supports common file types like CSV, Excel files, JSON, and Zeek network logs."),
        ("See the Attack Timeline:", "Shows every single event in order of time (what happened first, second, third)."),
        ("Spot Dangerous Activity:", "Automatically alerts you if someone is scanning ports, downloading viruses, or stealing data."),
        ("View an Interactive Map:", "Displays a visual diagram of all computers talking to each other (bad computers turn red)."),
        ("Download a PDF Report:", "With 1 click, creates a formal PDF incident report ready to print or email.")
    ])

    set_note(s2,
        "To explain the big picture to your teacher: "
        "Think of this project as a CCTV camera playback system, but for computer networks. "
        "Usually, network logs are just massive text files with millions of confusing numbers. "
        "Our project takes those text files and turns them into a clean story with charts, an interactive timeline, "
        "and a visual map. It lets an administrator see who got hacked, when it started, and what data was stolen in just a few clicks."
    )

    # =========================================================================
    # SLIDE 3: WHY IS CAMPUS NETWORK SECURITY HARD?
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "Why is Campus Network Security So Difficult?")

    cards_s3 = [
        (Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.3), "1. Thousands of Unknown Devices", RED, [
            ("Laptops & Phones Everywhere:", "Students, teachers, and guests bring their own laptops and phones onto campus Wi-Fi."),
            ("No Control:", "The college IT team cannot control what software or games students install on their personal devices."),
            ("Accidental Infection:", "One infected laptop can silently attack college lab servers from inside the Wi-Fi.")
        ]),
        (Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.3), "2. IP Addresses Keep Changing", AMBER, [
            ("Dynamic DHCP:", "When a student walks from the library to the canteen, their laptop disconnects and gets a brand new IP address."),
            ("Confusion during Investigation:", "An IP address that did an attack at 9:00 AM might belong to a totally innocent student by 10:00 AM!"),
            ("Hard to Track:", "Tracking who owned which IP at which minute is very difficult without our tool.")
        ]),
        (Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.3), "3. Millions of Log Lines", CYAN, [
            ("Information Overload:", "A normal campus produces over 100,000 log entries every single hour."),
            ("Humanly Impossible:", "No human can manually read text logs fast enough during a live cyber attack."),
            ("Delayed Action:", "By the time the IT team notices, the attacker has already stolen exam papers or sensitive student data.")
        ])
    ]

    for (cl, ct, cw, ch, ctitle, ccol, cpts) in cards_s3:
        add_card(s3, cl, ct, cw, ch, ctitle, ccol)
        add_bullet_points(s3, cl + Inches(0.2), ct + Inches(0.6), cw - Inches(0.4), ch - Inches(0.7), cpts)

    set_note(s3,
        "Here are the 3 real-world problems we solve: "
        "First, college Wi-Fi is crowded with thousands of personal student phones and laptops that might have viruses. "
        "Second, IP addresses change constantly throughout the day via DHCP. If an attack happens, it is very hard to tell who had that IP at that exact second. "
        "Third, there are simply too many log lines for a human to read. Our project automates this entire process."
    )

    # =========================================================================
    # SLIDE 4: PROJECT GOALS & OBJECTIVES
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "Project Goals: What We Built")

    add_card(s4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3), "Key Goals We Set Out to Achieve", CYAN)
    add_bullet_points(s4, Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.5), [
        ("Goal 1: Read Any Log File:", "Make a system that can accept logs from different firewalls, Wi-Fi routers, and web servers without crashing."),
        ("Goal 2: Fix Messy Time & Dates:", "Convert all different time formats (like Unix seconds, AM/PM, and dates) into one unified clock."),
        ("Goal 3: Catch Attacks Automatically:", "Build smart rules to detect hacking behavior like port scanning, brute force, and data theft."),
        ("Goal 4: Link Related Events:", "Connect the dots so the analyst sees the full chain of events, not just isolated alerts.")
    ])

    add_card(s4, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3), "User Experience & Presentation Goals", GREEN)
    add_bullet_points(s4, Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.5), [
        ("Goal 5: Simple Visual Dashboard:", "Build a beautiful Dark Mode interface that looks like a real-world Security Operations Center (SOC)."),
        ("Goal 6: Interactive Network Graph:", "Show devices as bubbles and connections as lines, so anyone can see the attack visually."),
        ("Goal 7: Save Cases & Findings:", "Let the security officer create cases, bookmark important evidence, and write down notes."),
        ("Goal 8: Create Professional Reports:", "Export clean PDF reports with 1 click so they can be submitted to college management.")
    ])

    set_note(s4,
        "Teacher, our goals were simple: "
        "1. Ingest files easily. "
        "2. Fix all timestamps so they line up on a single timeline. "
        "3. Automatically detect suspicious attacks. "
        "4. Present everything in a clear visual dashboard with an interactive map, and allow 1-click PDF export. "
        "Every single one of these goals has been fully built and tested in this project."
    )

    # =========================================================================
    # SLIDE 5: HOW THE SYSTEM WORKS (4 SIMPLE STEPS)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "How the System Works: 4 Simple Steps")

    steps = [
        ("Step 1: Upload", "Ingest Log Files", [
            ("User Uploads Logs:", "Upload CSV, JSON, or Zeek log files from Wi-Fi routers and firewalls."),
            ("Automatic Check:", "The system checks file format and columns automatically.")
        ], Inches(0.8), CYAN),
        ("Step 2: Clean", "Normalize the Data", [
            ("Standardize Clocks:", "Converts all dates and times into one matching time zone."),
            ("Clean IP & Ports:", "Verifies IP numbers and labels protocols (like HTTP, DNS, SSH).")
        ], Inches(3.8), BLUE_PILL),
        ("Step 3: Detect", "Run Security Rules", [
            ("Rule Engine:", "Checks for port scans, backdoor ports, and rapid data uploads."),
            ("Connect the Dots:", "Links which computer contacted which server and when.")
        ], Inches(6.8), AMBER),
        ("Step 4: Show", "Visuals & Reports", [
            ("Interactive Map:", "Draws a live network graph of all computers."),
            ("1-Click PDF:", "Generates a clean PDF incident report.")
        ], Inches(9.8), GREEN)
    ]

    for (stitle, ssub, sitems, sx, scol) in steps:
        add_card(s5, sx, Inches(1.6), Inches(2.7), Inches(5.3), stitle, scol)
        subb = s5.shapes.add_textbox(sx + Inches(0.15), Inches(2.1), Inches(2.4), Inches(0.35))
        stf = subb.text_frame
        stf.margin_left = stf.margin_top = 0
        sp = stf.paragraphs[0]
        sp.text = ssub.upper()
        sp.font.name = "Arial"
        sp.font.size = Pt(9)
        sp.font.bold = True
        sp.font.color.rgb = MUTED

        add_bullet_points(s5, sx + Inches(0.15), Inches(2.55), Inches(2.4), Inches(4.2), sitems)

    set_note(s5,
        "Here is the simple step-by-step workflow: "
        "Step 1: The user uploads raw log files. "
        "Step 2: The cleaning engine standardizes the timestamps and IP addresses into one format. "
        "Step 3: Our detection rules inspect the events to find suspicious hacking patterns. "
        "Step 4: The results are displayed on the web dashboard with charts, an interactive graph, and a PDF report generator."
    )

    # =========================================================================
    # SLIDE 6: TECHNOLOGIES USED & WHY WE CHOSE THEM
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, "Technologies Used & Why We Chose Them")

    techs = [
        ("Python 3 (Backend)", [
            ("Why Python?", "Python is the #1 programming language for cybersecurity and data handling."),
            ("Flask Web Framework:", "Fast, lightweight web framework that connects the backend logic to the web pages."),
            ("Pandas:", "Used for quick tabular data filtering and CSV exports.")
        ], Inches(0.8), Inches(1.6), Inches(5.6), Inches(2.55), CYAN),
        ("SQLite (Database)", [
            ("Why SQLite?", "Requires zero installation or setup. The database is a single file inside the project."),
            ("Fast Query Speed:", "We added special indexes on timestamps and IP addresses so searches take less than 0.1 second."),
            ("Private & Secure:", "All data stays on the local computer. No cloud service is required.")
        ], Inches(0.8), Inches(4.35), Inches(5.6), Inches(2.55), GREEN),
        ("Bootstrap 5 & Chart.js (UI)", [
            ("Bootstrap 5 Dark Mode:", "Gives a sleek, professional Security Operations Center (SOC) look."),
            ("Chart.js:", "Draws animated bar charts and donut charts that show peak traffic hours."),
            ("Responsive Design:", "Works smoothly on laptops, tablets, and external monitors.")
        ], Inches(6.8), Inches(1.6), Inches(5.7), Inches(2.55), AMBER),
        ("Cytoscape.js & ReportLab", [
            ("Cytoscape.js (Network Graph):", "A powerful library that draws circles for computers and lines for connections."),
            ("ReportLab (PDF Generator):", "Builds clean, printable PDF reports with tables and headers directly from Python code.")
        ], Inches(6.8), Inches(4.35), Inches(5.7), Inches(2.55), CYAN)
    ]

    for (ctitle, citems, cx, cy, cw, ch, ccol) in techs:
        add_card(s6, cx, cy, cw, ch, ctitle, ccol)
        add_bullet_points(s6, cx + Inches(0.2), cy + Inches(0.55), cw - Inches(0.4), ch - Inches(0.6), citems)

    set_note(s6,
        "Here is our technology stack: "
        "We chose Python and Flask because they are lightweight, reliable, and standard in cybersecurity. "
        "We chose SQLite because it needs zero configuration and keeps all sensitive college logs completely private on the local machine. "
        "For visualization, we used Chart.js for graphs and Cytoscape.js to build the interactive network map. "
        "Finally, ReportLab creates professional PDF reports directly."
    )

    # =========================================================================
    # SLIDE 7: STEP 1 - READING & CLEANING LOGS
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7)
    add_header(s7, "Step 1: Reading & Cleaning Messy Logs")

    add_card(s7, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3), "Handling Different Log Formats", CYAN)
    add_bullet_points(s7, Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.5), [
        ("Accepts Many File Types:", "Can read CSV files, TSV files, JSON records, and Zeek network security logs."),
        ("Understands Different Column Names:", "If one file calls an IP 'Source_IP' and another calls it 'src_addr' or 'id.orig_h', our system automatically recognizes that they mean the same thing!"),
        ("Safe with Missing Data:", "If some fields like hostname or description are missing, the system keeps working smoothly without crashing."),
        ("Keeps the Original Data:", "Saves the raw original log line so security analysts can always double-check the original proof.")
    ])

    add_card(s7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3), "Fixing Dates & IP Addresses", GREEN)
    add_bullet_points(s7, Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.5), [
        ("Aligning the Clocks:", "Converts all dates (Unix epoch numbers, AM/PM, standard dates) into one unified international time format."),
        ("Checking IP Validity:", "Ensures every IP address is a real IPv4 address (e.g. 10.0.1.45) and removes garbage entries."),
        ("Labeling Event Types:", "Automatically tags events as Web (HTTP/TLS), DNS, Email (SMTP), or DHCP based on ports and protocols."),
        ("Color-Coded Severity:", "Gives each event a clear severity badge: Informational, Low, Medium, High, or Critical.")
    ])

    set_note(s7,
        "This slide explains how we clean up messy logs: "
        "In real life, a Zeek log calls the IP 'id.orig_h', while a firewall calls it 'src_ip'. "
        "Our code automatically recognizes that these mean the same thing. "
        "More importantly, it puts all timestamps on the same clock so events can be sorted in exact order. "
        "Without this step, it is impossible to know what happened first during an attack."
    )

    # =========================================================================
    # SLIDE 8: STEP 2 - CONNECTING THE DOTS (CORRELATION)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8)
    add_header(s8, "Step 2: Connecting the Dots (Correlation Engine)")

    corr_items = [
        ("1. Shared Computer Tracking", [
            ("Tracking Connections:", "Finds all computers that spoke to the same target server or were spoken to by a hacker."),
            ("Pivot Search:", "Clicking any computer instantly shows all its incoming and outgoing connections.")
        ], Inches(0.8), Inches(1.6), Inches(5.6), Inches(2.55), CYAN),
        ("2. Following the DNS Clues", [
            ("Website Name to IP:", "When a computer looks up a bad website (like malware-c2.xyz), it immediately tries to connect to that website's IP address."),
            ("Connecting the Link:", "Our system links the DNS query directly to the following connection attempt.")
        ], Inches(6.8), Inches(1.6), Inches(5.7), Inches(2.55), AMBER),
        ("3. Time Window Clustering", [
            ("Group Events in 5-Minute Blocks:", "Looks at events happening close together in time."),
            ("Spots Fast Bursts:", "Helps identify sudden bursts of activity, like an attacker quickly stealing data within 2 minutes.")
        ], Inches(0.8), Inches(4.35), Inches(5.6), Inches(2.55), GREEN),
        ("4. Related Event Finder", [
            ("Instant Related Events:", "When an analyst clicks any event, the system shows all other events connected to it."),
            ("Clear Context:", "Analyst sees the whole story instead of just one isolated alert.")
        ], Inches(6.8), Inches(4.35), Inches(5.7), Inches(2.55), BLUE_PILL)
    ]

    for (ctitle, citems, cx, cy, cw, ch, ccol) in corr_items:
        add_card(s8, cx, cy, cw, ch, ctitle, ccol)
        add_bullet_points(s8, cx + Inches(0.2), cy + Inches(0.55), cw - Inches(0.4), ch - Inches(0.6), citems)

    set_note(s8,
        "Here is how our Correlation Engine connects the dots: "
        "A single log line in isolation looks harmless. For example, visiting a website is normal. "
        "But if a computer looks up a suspicious website name, and two seconds later opens a hidden connection to that server, "
        "that is an attack! Our correlation engine links these clues together so the analyst sees the complete picture."
    )

    # =========================================================================
    # SLIDE 9: STEP 3 - THE 6 CYBER ATTACKS WE CATCH
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_bg(s9)
    add_header(s9, "Step 3: The 6 Common Cyber Attacks We Catch")

    table_s9 = s9.shapes.add_table(7, 3, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.2))
    tbl = table_s9.table
    tbl.columns[0].width = Inches(3.0)
    tbl.columns[1].width = Inches(6.733)
    tbl.columns[2].width = Inches(2.0)

    t_headers = ["Attack Type", "What Does the Attacker Do? (In Plain English)", "Severity"]
    for i, h in enumerate(t_headers):
        cell = tbl.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(30, 41, 59)
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = CYAN

    attack_rows = [
        ("1. Port Scanning", "Checking many doors (ports) on a single server to find an open door to break in.", "HIGH"),
        ("2. Network Sweep", "Scanning across multiple computers on campus to see which ones are online.", "HIGH"),
        ("3. Suspicious DNS Lookups", "Querying dangerous domain names (like .xyz, .tk, .top) used by malware.", "HIGH"),
        ("4. Unusual / Backdoor Port", "Connecting to known hacker ports (like 4444 used by Metasploit) to control a PC.", "CRITICAL"),
        ("5. Rapid Traffic Burst", "Uploading huge amounts of data in a short time (data theft / exfiltration).", "HIGH"),
        ("6. Password Brute Force", "Trying to guess passwords repeatedly (like 15 failed SSH logins in a row).", "MEDIUM")
    ]

    for row_idx, r in enumerate(attack_rows, start=1):
        for col_idx, val in enumerate(r):
            cell = tbl.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if row_idx % 2 == 1 else RGBColor(20, 27, 42)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = "Arial"
            p.font.size = Pt(11)
            if col_idx == 0:
                p.font.bold = True
                p.font.color.rgb = WHITE
            elif col_idx == 2:
                p.font.bold = True
                p.font.color.rgb = RED if "CRITICAL" in val or "HIGH" in val else AMBER
            else:
                p.font.color.rgb = LIGHT_GRAY

    set_note(s9,
        "Teacher, these are the 6 attacks our system catches automatically. "
        "1. Port Scanning: checking doors on a server. "
        "2. Network Sweep: finding other computers on the Wi-Fi. "
        "3. Suspicious DNS: looking up dangerous websites. "
        "4. Unusual Ports: hacker backdoors like port 4444. "
        "5. Traffic Burst: stealing data rapidly. "
        "6. Brute Force: guessing passwords over and over. "
        "Every rule is clear, explainable, and stores exact proof of what happened."
    )

    # =========================================================================
    # SLIDE 10: THE DEMO ATTACK STORY (STEP-BY-STEP)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_bg(s10)
    add_header(s10, "The Demo Attack Story (What We Simulated)")

    demo_stages = [
        ("Phase 1: Student Connects (08:00 AM)", [
            ("Normal Arrival:", "Student laptop 10.0.1.45 connects to campus Wi-Fi and gets an IP via DHCP."),
            ("Normal Browsing:", "Visits normal college portals (portal.campus.edu and moodle) to look innocent.")
        ], Inches(0.8), Inches(1.6), Inches(5.6), Inches(2.55), BLUE_PILL),
        ("Phase 2: Scanning the Server (09:00 AM)", [
            ("Probing Doors:", "The laptop starts scanning the Campus Web Server (10.0.2.10) on 10 different ports."),
            ("First Alert:", "Our system flags this as 'Port Scanning' within seconds.")
        ], Inches(6.8), Inches(1.6), Inches(5.7), Inches(2.55), AMBER),
        ("Phase 3: Backdoor Shell (09:01 AM)", [
            ("Dangerous Website:", "Queries malware-c2.xyz to contact the attacker's server."),
            ("Hacker Takes Control:", "Connects to port 4444 (Metasploit shell) — flagged as CRITICAL.")
        ], Inches(0.8), Inches(4.35), Inches(5.6), Inches(2.55), RED),
        ("Phase 4: Stealing Data & Moving (09:02 AM)", [
            ("Data Exfiltration:", "Uploads 20 chunks of sensitive data in 2 minutes (Traffic Burst alert)."),
            ("Second Computer (10.0.1.78):", "Tries to break into admin servers using 15 failed SSH password attempts.")
        ], Inches(6.8), Inches(4.35), Inches(5.7), Inches(2.55), CYAN)
    ]

    for (ctitle, citems, cx, cy, cw, ch, ccol) in demo_stages:
        add_card(s10, cx, cy, cw, ch, ctitle, ccol)
        add_bullet_points(s10, cx + Inches(0.2), cy + Inches(0.55), cw - Inches(0.4), ch - Inches(0.6), citems)

    set_note(s10,
        "To test our system, we generated a realistic campus attack story with ~400 events. "
        "At 8:00 AM, a student laptop connects to Wi-Fi normally. "
        "At 9:00 AM, it suddenly scans the campus web server. "
        "At 9:01 AM, it opens a backdoor connection to a hacker server on port 4444. "
        "At 9:02 AM, it rapidly uploads stolen files. Then another laptop tries guessing admin passwords. "
        "Our system detects and reconstructs every single phase of this attack."
    )

    # =========================================================================
    # SLIDE 11: THE MAIN DASHBOARD
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_bg(s11)
    add_header(s11, "The Main Dashboard: At-a-Glance Control")

    stat_cards = [
        ("TOTAL LOG EVENTS", "~400 Events", "Cleaned & Indexed", CYAN),
        ("ACTIVE COMPUTERS", "35+ Devices", "Laptops, Servers, Wi-Fi", GREEN),
        ("SECURITY ALERTS", "10+ Warnings", "Detected Automatically", RED),
        ("ALERT SEVERITY", "Critical / High", "Prioritized for Action", AMBER)
    ]
    for i, (lbl, val, sub, col) in enumerate(stat_cards):
        mx = Inches(0.8 + i * 2.98)
        mc = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, mx, Inches(1.6), Inches(2.8), Inches(1.2))
        mc.fill.solid()
        mc.fill.fore_color.rgb = CARD_BG
        mc.line.color.rgb = CARD_BORDER
        tf = mc.text_frame
        tf.margin_left = Inches(0.15)
        tf.margin_top = Inches(0.12)
        p1 = tf.paragraphs[0]
        p1.text = lbl
        p1.font.name = "Arial"
        p1.font.size = Pt(8.5)
        p1.font.bold = True
        p1.font.color.rgb = col
        p2 = tf.add_paragraph()
        p2.text = val
        p2.font.name = "Arial"
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p3 = tf.add_paragraph()
        p3.text = sub
        p3.font.name = "Arial"
        p3.font.size = Pt(8.5)
        p3.font.color.rgb = MUTED

    add_card(s11, Inches(0.8), Inches(3.05), Inches(5.6), Inches(3.85), "Visual Charts for Fast Understanding", CYAN)
    add_bullet_points(s11, Inches(1.0), Inches(3.55), Inches(5.2), Inches(3.2), [
        ("Activity Timeline Chart:", "A bar chart that shows network traffic hour-by-hour. You can immediately see the big spike at 9:00 AM when the attack started!"),
        ("Protocol Donut Chart:", "Shows a breakdown of what kinds of traffic occurred (Web, DNS, Email, SSH).")
    ])

    add_card(s11, Inches(6.8), Inches(3.05), Inches(5.7), Inches(3.85), "Immediate Action Panels", AMBER)
    add_bullet_points(s11, Inches(7.0), Inches(3.55), Inches(5.3), Inches(3.2), [
        ("Top Talkers List:", "Shows which IP addresses sent the most traffic. The hacker's laptop is right at the top!"),
        ("Live Alerts Feed:", "Lists every detected threat with its severity badge. Clicking an alert lets you jump straight to the evidence.")
    ])

    set_note(s11,
        "This is the Main Dashboard. When an administrator opens the website, this is what they see. "
        "At the top, 4 stat cards summarize the network: 400 events, 35 hosts, and 10 alerts. "
        "The bar chart shows a clear spike at 9:00 AM, making it obvious when the attack occurred. "
        "The Top Talkers list immediately exposes the suspicious laptop."
    )

    # =========================================================================
    # SLIDE 12: THE ATTACK TIMELINE VIEW
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_bg(s12)
    add_header(s12, "The Attack Timeline: Step-by-Step History")

    add_card(s12, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3), "How the Timeline Helps the Investigator", CYAN)
    add_bullet_points(s12, Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.5), [
        ("Strict Order of Time:", "Every event is arranged from earliest to latest, like chapters in a book."),
        ("Easy Search & Filters:", "Filter by IP address (e.g. type '10.0.1.45'), protocol (HTTP or TCP), or severity level."),
        ("Color-Coded Badges:", "Green for normal, Yellow for medium risk, Red for critical attacks."),
        ("Instant Results:", "Filtering through hundreds of records happens instantly without refreshing the page.")
    ])

    add_card(s12, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3), "Click Any Event to See Proof", GREEN)
    add_bullet_points(s12, Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.5), [
        ("Full Detail Pop-Up:", "Clicking any event opens a detailed pop-up window."),
        ("Exact Details:", "Shows exact time down to the second, source port, destination port, and device name."),
        ("Raw Data Proof:", "Shows the raw text line from the original log file for legal proof."),
        ("One-Click Pin to Case:", "Allows the analyst to bookmark this event into an active investigation case.")
    ])

    set_note(s12,
        "Slide 12 covers the Timeline view. "
        "Time is the most important element in an investigation. "
        "Our timeline arranges every event in order. If an analyst enters '10.0.1.45', the screen filters down to show only that device's actions. "
        "Clicking any event opens a window showing the exact time, ports, and raw proof. "
        "This turns hours of manual log reading into seconds of point-and-click investigation."
    )

    # =========================================================================
    # SLIDE 13: INTERACTIVE NETWORK MAP
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_bg(s13)
    add_header(s13, "Interactive Network Map (Cytoscape.js)")

    add_card(s13, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3), "A Visual Map of All Connected Computers", CYAN)
    add_bullet_points(s13, Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.5), [
        ("Circles & Lines:", "Each computer or server is drawn as a circle (node). Connections between them are drawn as lines (edges)."),
        ("Clear Color Meaning:", "Normal computers are BLUE. Servers are YELLOW. Infected/attacker computers glow RED!"),
        ("Drag & Explore:", "You can zoom in, zoom out, drag computers around, and rearrange the view with your mouse."),
        ("Different Layout Styles:", "Switch between physics layout, circle layout, or grid layout with one button.")
    ])

    add_card(s13, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3), "Click to Inspect Any Computer", GREEN)
    add_bullet_points(s13, Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.5), [
        ("Side Information Drawer:", "Clicking any computer bubble opens a panel showing its IP address, hostname, and role."),
        ("See Who It Talked To:", "Shows total connections sent and received, and any alerts linked to that device."),
        ("Exposes the Attack Path:", "You can visually trace the line from the red infected laptop to the external hacker server!"),
        ("Great for Non-Technical Leaders:", "Makes it easy to explain an incident to non-technical college principals or deans.")
    ])

    set_note(s13,
        "Slide 13 shows our interactive Network Map. "
        "While the timeline tells you 'when' an attack happened, this map shows you 'who talked to whom'. "
        "Normal computers appear in blue, servers in yellow, and bad computers glow red. "
        "You can drag nodes around and click any computer to see its connections. "
        "This makes it extremely easy to explain an attack to a teacher, dean, or police officer."
    )

    # =========================================================================
    # SLIDE 14: INVESTIGATION CASE WORKSPACE
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_bg(s14)
    add_header(s14, "Investigation Workspace: Managing Cases")

    inv_items = [
        ("1. Create an Investigation Case", [
            ("Organized Cases:", "Create a named case like 'Library Wi-Fi Incident - Oct 15' with an assigned analyst name."),
            ("Time Framing:", "Set a specific time window to focus only on the incident hours.")
        ], Inches(0.8), Inches(1.6), Inches(5.6), Inches(2.55), CYAN),
        ("2. Pinning Evidence", [
            ("Bookmark Evidence:", "Pin suspicious events and suspect computer IPs directly into the case folder."),
            ("Clean Sub-Timeline:", "View an isolated timeline containing only the bookmarked evidence.")
        ], Inches(6.8), Inches(1.6), Inches(5.7), Inches(2.55), AMBER),
        ("3. Writing Analyst Notes", [
            ("Document Findings:", "Write down what was discovered, what action was taken, and recommendations."),
            ("Official Record:", "Notes are permanently saved and included in the final report.")
        ], Inches(0.8), Inches(4.35), Inches(5.6), Inches(2.55), GREEN),
        ("4. Alert Triage Lifecycle", [
            ("Simple Status Buttons:", "Move alerts through clear stages: Open -> Acknowledged -> Resolved -> False Positive."),
            ("No Confusion:", "Ensures team members don't duplicate work on the same incident.")
        ], Inches(6.8), Inches(4.35), Inches(5.7), Inches(2.55), BLUE_PILL)
    ]

    for (ctitle, citems, cx, cy, cw, ch, ccol) in inv_items:
        add_card(s14, cx, cy, cw, ch, ctitle, ccol)
        add_bullet_points(s14, cx + Inches(0.2), cy + Inches(0.55), cw - Inches(0.4), ch - Inches(0.6), citems)

    set_note(s14,
        "Finding an attack is only half the job; documenting it is the other half. "
        "In our Investigation Workspace, an analyst can open a formal case, bookmark suspicious events, "
        "tag suspect computers, and write down notes. "
        "Alerts can be marked as Resolved or False Positive. This gives a structured workflow just like professional cybersecurity tools."
    )

    # =========================================================================
    # SLIDE 15: ONE-CLICK PDF REPORT GENERATION
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    set_bg(s15)
    add_header(s15, "1-Click PDF Incident Report Generator")

    add_card(s15, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3), "What is in the PDF Report?", CYAN)
    add_bullet_points(s15, Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.5), [
        ("Executive Summary:", "A high-level summary explaining what happened, how many computers were affected, and the risk level."),
        ("Incident Metadata:", "Date, time range, analyst name, and official confidentiality header."),
        ("Formatted Timeline Table:", "A neat table listing all attack events in order of time with source IP, destination IP, port, and protocol."),
        ("Triggered Alerts Breakdown:", "Lists all security rules that were triggered along with the technical evidence."),
        ("Analyst Notes & Recommendations:", "Includes recommendations on how to block the attacker and secure the network.")
    ])

    add_card(s15, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3), "Data Exports (CSV & JSON)", GREEN)
    add_bullet_points(s15, Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.5), [
        ("Instant PDF Download:", "Built using Python ReportLab. Generates a clean, printable PDF in under 2 seconds."),
        ("CSV Export:", "Download events into an Excel-compatible CSV file for offline viewing or printing."),
        ("JSON Export:", "Export clean JSON data to feed into external tools like Splunk or Elasticsearch."),
        ("No Third-Party Cloud Needed:", "The PDF is created locally on your computer with complete privacy.")
    ])

    set_note(s15,
        "Slide 15 highlights our reporting engine. "
        "When an investigation is finished, management needs a written report. "
        "With one click on the 'Export PDF' button, our Python ReportLab engine compiles an executive-ready PDF report. "
        "It includes a summary, chronological table, alert evidence, and analyst notes. "
        "Users can also export the data as CSV or JSON."
    )

    # =========================================================================
    # SLIDE 16: TESTING & PROOF IT WORKS
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    set_bg(s16)
    add_header(s16, "Testing & Proof: 100% Passed (68 Tests)")

    add_card(s16, Inches(0.8), Inches(1.6), Inches(11.733), Inches(1.2), "Automated Pytest Results: 68 / 68 Tests Passed Successfully", GREEN)
    tbox_test = s16.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(11.3), Inches(0.6))
    ttf = tbox_test.text_frame
    tp = ttf.paragraphs[0]
    tp.text = "We wrote 68 automated software tests. Every single test passed in under 13 seconds with ZERO errors."
    tp.font.name = "Arial"
    tp.font.size = Pt(12)
    tp.font.bold = True
    tp.font.color.rgb = WHITE

    test_cols = [
        ("File Parsers", "22 Tests Passed", [
            ("CSV & JSON Parsing:", "Tested with broken headers, missing data, and empty files."),
            ("Time Cleaning:", "Tested Unix timestamps and standard dates."),
            ("IP Validation:", "Verified real vs invalid IP addresses.")
        ], Inches(0.8), CYAN),
        ("Timeline & Linking", "14 Tests Passed", [
            ("Time Sorting:", "Tested strict chronological sorting."),
            ("Port Scan Links:", "Confirmed scanning events link together."),
            ("DNS Linking:", "Verified domain lookups link to IP handshakes.")
        ], Inches(3.8), AMBER),
        ("Attack Detectors", "9 Tests Passed", [
            ("6 Threat Rules:", "Tested all 6 rules with real and synthetic data."),
            ("Custom Limits:", "Tested changing settings in the GUI."),
            ("Zero False Alarms:", "Tested that normal browsing does NOT trigger alerts.")
        ], Inches(6.8), RED),
        ("APIs & Reports", "23 Tests Passed", [
            ("All Web Routes:", "Tested dashboard, upload, and timeline endpoints."),
            ("PDF Generation:", "Verified PDF report compiles without crashing."),
            ("CSV Exports:", "Checked tabular exports.")
        ], Inches(9.8), BLUE_PILL)
    ]

    for (stitle, scount, sitems, sx, scol) in test_cols:
        add_card(s16, sx, Inches(3.0), Inches(2.7), Inches(3.9), stitle, scol)
        cbox = s16.shapes.add_textbox(sx + Inches(0.15), Inches(3.45), Inches(2.4), Inches(0.3))
        ctf = cbox.text_frame
        ctf.margin_left = ctf.margin_top = 0
        cp = ctf.paragraphs[0]
        cp.text = scount.upper()
        cp.font.name = "Arial"
        cp.font.size = Pt(9)
        cp.font.bold = True
        cp.font.color.rgb = GREEN
        add_bullet_points(s16, sx + Inches(0.15), Inches(3.8), Inches(2.4), Inches(3.0), sitems)

    set_note(s16,
        "Teacher, we did not just write the application; we thoroughly tested it. "
        "We wrote an automated test suite of 68 individual test cases using pytest. "
        "All 68 tests passed with 100% success in under 13 seconds. "
        "This proves that our parsers don't crash on bad files, our detection rules don't produce false alarms, "
        "and our PDF reports generate smoothly every single time."
    )

    # =========================================================================
    # SLIDE 17: SUMMARY, FUTURE PLANS & THANK YOU
    # =========================================================================
    s17 = prs.slides.add_slide(blank_layout)
    set_bg(s17)
    add_header(s17, "Project Summary, Future Plans & Q&A")

    add_card(s17, Inches(0.8), Inches(1.6), Inches(5.6), Inches(3.6), "What We Plan to Add Next (Future Work)", CYAN)
    add_bullet_points(s17, Inches(1.0), Inches(2.2), Inches(5.2), Inches(2.8), [
        ("Machine Learning (AI):", "Add AI algorithms to automatically detect brand new, unknown attack patterns without writing rules."),
        ("Live Packet Sniffing:", "Connect directly to campus Wi-Fi hardware to capture and analyze packets live."),
        ("Threat Intelligence:", "Automatically check unknown IP addresses against global blacklists (VirusTotal / AbuseIPDB).")
    ])

    add_card(s17, Inches(6.8), Inches(1.6), Inches(5.7), Inches(3.6), "Key Takeaways (Conclusion)", GREEN)
    add_bullet_points(s17, Inches(7.0), Inches(2.2), Inches(5.3), Inches(2.8), [
        ("Fully Working Solution:", "Delivered a complete, working platform for campus network incident response."),
        ("Saves Time:", "Reduces manual log reading from days down to a 3-minute visual investigation."),
        ("100% Offline & Private:", "Safe for college use because student data never leaves the campus.")
    ])

    # Big Q&A Box
    qa_card = s17.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.4), Inches(11.733), Inches(1.5))
    qa_card.fill.solid()
    qa_card.fill.fore_color.rgb = RGBColor(16, 23, 40)
    qa_card.line.color.rgb = RGBColor(37, 99, 235)
    qa_card.line.width = Pt(1.5)

    qatf = qa_card.text_frame
    qatf.margin_top = Inches(0.2)
    qap1 = qatf.paragraphs[0]
    qap1.alignment = PP_ALIGN.CENTER
    qap1.text = "Thank You for Your Time & Guidance"
    qap1.font.name = "Arial"
    qap1.font.size = Pt(20)
    qap1.font.bold = True
    qap1.font.color.rgb = WHITE

    qap2 = qatf.add_paragraph()
    qap2.alignment = PP_ALIGN.CENTER
    qap2.text = "Questions & Answers | Ready for Live Project Demonstration"
    qap2.font.name = "Arial"
    qap2.font.size = Pt(13)
    qap2.font.bold = True
    qap2.font.color.rgb = CYAN

    set_note(s17,
        "In conclusion, our project solves a major real-world headache for college IT teams. "
        "It turns chaotic network logs into a simple timeline and visual map, making it fast and easy to catch attackers. "
        "In the future, we hope to add machine learning for zero-day threats and live packet capture. "
        "Thank you very much, teacher. I am now ready for your questions and to show the live demo."
    )

    output_path = os.path.abspath("Campus_Network_Attack_Timeline_Reconstructor_Presentation.pptx")
    prs.save(output_path)
    print(f"Successfully generated 17-slide simplified presentation at: {output_path}")
    return output_path

if __name__ == "__main__":
    create_simple_deck()
